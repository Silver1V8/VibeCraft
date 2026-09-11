"""Excel → Markdown per l'intake: lettura DIRETTA dell'originale, XLSX con openpyxl e XLS con
xlrd, senza normalizzazioni intermedie né tabelle dedotte. Prove e scelte:
rag/eval/results/excel_diretto_20260911/RESOCONTO.md.

Contratto della rappresentazione:
- una sezione `##` per foglio, nell'ordine del workbook, compresi i fogli vuoti o nascosti;
- ogni tabella Markdown è un intervallo di celle `### A1:F5`: prima colonna il numero di riga,
  intestazioni = lettere di colonna (coordinate, mai intestazioni dei dati);
- una colonna senza valori né unioni in TUTTO il foglio separa gli intervalli; dentro un
  intervallo di colonne, una riga vuota separa i blocchi. Regola solo geometrica: niente
  parole chiave, soglie o coordinate speciali. Due tabelle affiancate da una colonna vuota
  restano distinte; una riga con celle vuote intermedie resta una riga, allineata alle colonne
  del suo intervallo (anche dopo una riga vuota);
- il valore MEMORIZZATO: formule mai ricalcolate, numeri senza formato di visualizzazione (i
  formati non generali sono elencati per foglio), date e ore ISO al millisecondo, booleani
  TRUE/FALSE, errori col codice Excel; il testo è verbatim a meno degli escape Markdown, e
  quello che sembrerebbe un altro tipo è elencato per foglio;
- celle unite: valore solo nella cella iniziale, intervallo elencato.
Niente troncamenti: oltre MAX_CELLS celle da rappresentare l'import è rifiutato."""
from __future__ import annotations

import io
import re
import struct
import warnings
from bisect import bisect_right
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, time, timedelta
from xml.etree.ElementTree import iterparse
from zipfile import ZipFile

from openpyxl.styles.numbers import is_date_format, is_datetime, is_timedelta_format
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel

from core.intake_convert import ConversionError

# Celle rappresentate (vuote comprese) per workbook. Oltre: errore esplicito, mai troncamento.
MAX_CELLS = 1_000_000
# Celle presenti nell'XML dei fogli XLSX (anche solo formattate), contate PRIMA di caricare.
# ponytail: due letture complete di openpyxl, misurate ~1,1 GB e 19 s a 990k celle; se servono
# file più grandi, leggere le formule in streaming invece della seconda lettura completa.
MAX_XML_CELLS = 2_000_000
# Somma delle parti ZIP PRIMA di decomprimere: include sharedStrings, stili e immagini,
# che il numero di celle non limita. Sopra 128 MiB errore esplicito, mai output parziale.
MAX_XLSX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024

LEGENDA = (
    "Contenuto delle celle del workbook, letto direttamente dal file. Ogni tabella è un "
    "intervallo di celle con le coordinate originali: la prima colonna è il numero di riga, "
    "A, B, C… sono le lettere di colonna, non intestazioni dei dati. Una colonna vuota in tutto "
    "il foglio separa gli intervalli e una riga vuota ne separa i blocchi: separazioni "
    "geometriche, non tabelle dichiarate nel file; celle con lo stesso numero di riga stanno "
    "sulla stessa riga del foglio. Valori memorizzati nel file: formule non ricalcolate, numeri "
    "senza formato di visualizzazione, date e ore ISO. Formattazione, convalide e oggetti "
    "grafici non sono riportati.")

# Punteggiatura che CommonMark/GFM interpreterebbe in una cella; `_` solo al bordo di una
# parola: dentro una parola è letterale e gli identificatori (codice_articolo) restano cercabili.
_MD_SPECIAL = re.compile(r"[\\`*\[\]<>|~&]|(?<![^\W_])_|_(?![^\W_])")
# Controlli C0/C1 (CR compreso) e separatori di riga Unicode come entità: nessun carattere
# spezza la riga della tabella e CRLF resta distinguibile da LF. LF diventa <br>, non ambiguo
# perché `<` dei dati è escapato.
_CONTROL = {cp: f"&#{cp};" for cp in (*range(0, 9), *range(11, 32), *range(127, 160), 0x2028, 0x2029)}
_EDGE_SPACE = re.compile(r"^\s+|\s+$")
# Testo che nel Markdown sembrerebbe un altro tipo: vuoto, numero, booleano, errore, data/ora.
_LOOKS_TYPED = re.compile(
    r"|[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?|TRUE|FALSE|#N/A|#[A-Z0-9/_]+[!?]"
    r"|\d{4}-\d{2}-\d{2}(?: \d{2}:\d{2}:\d{2}(?:\.\d+)?)?|-?\d+:\d{2}:\d{2}(?:\.\d+)?")
_STATO = {"hidden": "nascosto", "veryHidden": "molto nascosto", 1: "nascosto", 2: "molto nascosto"}


def _escape(text: str) -> str:
    text = _MD_SPECIAL.sub(lambda m: "\\" + m[0], text)
    text = text.translate(_CONTROL).replace("\n", "<br>")
    return _EDGE_SPACE.sub(lambda m: "".join(f"&#{ord(ch)};" for ch in m[0]), text)


def _number(value) -> str:
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e16:
        return str(int(value))
    return repr(value)  # la rappresentazione più breve che ridà esattamente il valore


def _temporal(value, fmt: str) -> str:
    if isinstance(value, timedelta):  # formati durata, es. [h]:mm:ss
        ms = round(value / timedelta(milliseconds=1))
        sign, ms = ("-" if ms < 0 else ""), abs(ms)
        s, ms = divmod(ms, 1000)
        m, s = divmod(s, 60)
        h, m = divmod(m, 60)
        return f"{sign}{h}:{m:02}:{s:02}" + (f".{ms:03}" if ms else "")
    spec = "milliseconds" if value.microsecond else "seconds"
    if isinstance(value, time):
        return value.isoformat(timespec=spec)
    # Solo la mezzanotte esatta con un formato di sola data perde l'ora: l'ora è 00:00:00.
    if is_datetime(fmt) == "date" and value.time() == time():
        return value.date().isoformat()
    return value.isoformat(sep=" ", timespec=spec)


def _ref(r1: int, c1: int, r2: int | None = None, c2: int | None = None) -> str:
    start = f"{get_column_letter(c1)}{r1}"
    return start if r2 is None or (r1, c1) == (r2, c2) else f"{start}:{get_column_letter(c2)}{r2}"


def _spans(pairs) -> list[list[int]]:
    """Intervalli [a, b] sovrapposti o adiacenti fusi."""
    out: list[list[int]] = []
    for a, b in sorted(pairs):
        if out and a <= out[-1][1] + 1:
            out[-1][1] = max(out[-1][1], b)
        else:
            out.append([a, b])
    return out


def _ranges(coords) -> str:
    """Celle (riga, colonna) → rettangoli compatti, es. «A3:A5, C7»."""
    by_row = defaultdict(list)
    for r, c in coords:
        by_row[r].append((c, c))
    rows_by_span = defaultdict(list)
    for r, cols in by_row.items():
        for c1, c2 in _spans(cols):
            rows_by_span[(c1, c2)].append((r, r))
    rects = [(r1, c1, r2, c2) for (c1, c2), rows in rows_by_span.items() for r1, r2 in _spans(rows)]
    return ", ".join(_ref(*rect) for rect in sorted(rects))


@dataclass
class _Sheet:
    name: str
    stato: str = ""
    celle: dict = field(default_factory=dict)          # (riga, colonna) -> testo Markdown
    unioni: list = field(default_factory=list)          # (r1, c1, r2, c2)
    formati: dict = field(default_factory=lambda: defaultdict(set))
    testo_tipizzato: set = field(default_factory=set)
    tabelle: list = field(default_factory=list)
    note: list = field(default_factory=list)            # (riga, colonna, Markdown): commenti, link

    def put(self, r: int, c: int, value, fmt: str = "General", error: bool = False) -> None:
        if error or isinstance(value, str):
            if not error and _LOOKS_TYPED.fullmatch(value):
                self.testo_tipizzato.add((r, c))
            text = _escape(value)
        elif isinstance(value, bool):
            text = "TRUE" if value else "FALSE"
        else:
            text = (_temporal(value, fmt) if isinstance(value, (datetime, time, timedelta))
                    else _number(value))
            if fmt and fmt != "General":
                self.formati[fmt].add((r, c))
        self.celle[(r, c)] = text

    def commento(self, r: int, c: int, autore: str | None, testo: str) -> None:
        chi = f" ({_escape(autore)})" if autore else ""
        self.note.append((r, c, f"Commento {_ref(r, c)}{chi}: {_escape(testo)}"))

    def collegamento(self, r: int, c: int, destinazione: str) -> None:
        self.note.append((r, c, f"Collegamento {_ref(r, c)}: {_escape(destinazione)}"))


def _blocks(s: _Sheet) -> list[tuple[int, int, int, int]]:
    bands = _spans([(c, c) for _, c in s.celle] + [(c1, c2) for _, c1, _, c2 in s.unioni])
    starts = [a for a, _ in bands]
    rows = defaultdict(list)
    for r, c in s.celle:
        rows[bisect_right(starts, c) - 1].append((r, r))
    for r1, c1, r2, _ in s.unioni:
        rows[bisect_right(starts, c1) - 1].append((r1, r2))
    return sorted((r1, bands[i][0], r2, bands[i][1]) for i, spans in rows.items()
                  for r1, r2 in _spans(spans))


def _render(sheets: list[_Sheet]) -> str:
    blocchi = [(s, _blocks(s)) for s in sheets]
    totale = sum((r2 - r1 + 1) * (c2 - c1 + 1) for _, bl in blocchi for r1, c1, r2, c2 in bl)
    if totale > MAX_CELLS:
        raise ConversionError(
            f"il workbook richiederebbe {totale:,} celle nel Markdown (limite {MAX_CELLS:,}): "
            f"intervalli troppo estesi. Suddividilo o togli i valori sparsi lontani dai dati, "
            f"poi ricaricalo".replace(",", "."))
    parts = [LEGENDA]
    for s, bl in blocchi:
        parts.append("## " + _escape(s.name).replace("#", "\\#"))
        info = [f"Foglio {s.stato}."] if s.stato else []
        if not s.celle:
            info.append("Nessuna cella con valore.")
        if s.unioni:
            info.append("Celle unite, valore solo nella prima cella: "
                        + ", ".join(_ref(*u) for u in sorted(s.unioni)) + ".")
        if s.formati:
            info.append("Formati numerici del file, non applicati ai valori mostrati: " + "; ".join(
                f"«{_escape(f)}» {_ranges(cc)}" for f, cc in sorted(s.formati.items())) + ".")
        if s.testo_tipizzato:
            info.append("Testo, non numero, data, booleano o errore: "
                        + _ranges(s.testo_tipizzato) + ".")
        if s.tabelle:
            info.append("Tabelle dichiarate nel file: "
                        + "; ".join(f"{_escape(n)} {ref}" for n, ref in s.tabelle) + ".")
        info += [testo for *_, testo in sorted(s.note)]
        if info:
            parts.append("\n".join("- " + i for i in info))
        for r1, c1, r2, c2 in bl:
            cols = range(c1, c2 + 1)
            righe = ["| Riga | " + " | ".join(get_column_letter(c) for c in cols) + " |",
                     "| --- " * (len(cols) + 1) + "|"]
            righe += ["| " + " | ".join([str(r), *(s.celle.get((r, c), "") for c in cols)]) + " |"
                      for r in range(r1, r2 + 1)]
            parts.append(f"### {_ref(r1, c1, r2, c2)}\n\n" + "\n".join(righe))
    return "\n\n".join(parts)


def _xml_cells(data: bytes) -> int:
    """Preflight streaming: anche le unioni allocano celle durante load_workbook."""
    n = merged = 0
    with ZipFile(io.BytesIO(data)) as zf:
        size = sum(part.file_size for part in zf.infolist())
        if size > MAX_XLSX_UNCOMPRESSED_BYTES:
            raise ConversionError(
                f"workbook decompresso troppo grande: {size} byte "
                f"(limite {MAX_XLSX_UNCOMPRESSED_BYTES} byte). Suddividilo e ricaricalo")
        for name in zf.namelist():
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                continue
            with zf.open(name) as fh:
                events = iterparse(fh, events=("start", "end"))
                _, root = next(events)
                for event, element in events:
                    if event == "start":
                        tag = element.tag.rsplit("}", 1)[-1]
                        if tag == "c":
                            n += 1
                            if n > MAX_XML_CELLS:
                                return n
                        elif tag == "mergeCell":
                            c1, r1, c2, r2 = range_boundaries(element.attrib["ref"])
                            merged += (r2 - r1 + 1) * (c2 - c1 + 1)
                            if merged > MAX_CELLS:
                                raise ConversionError(
                                    f"celle unite troppo estese (limite {MAX_CELLS} celle "
                                    f"per workbook): suddividi le unioni e ricaricalo")
                    else:
                        element.clear()
                        root.clear()
    return n


def _read_xlsx(data: bytes) -> tuple[list[_Sheet], list[str]]:
    import openpyxl  # noqa: PLC0415 (i test simulano un guasto sostituendo load_workbook)

    if (presenti := _xml_cells(data)) > MAX_XML_CELLS:
        raise ConversionError(
            f"il file contiene {presenti:,} celle nei fogli, comprese quelle solo formattate "
            f"(limite {MAX_XML_CELLS:,}): suddividilo o cancella la formattazione residua, poi "
            f"ricaricalo".replace(",", "."))

    # Due letture dello stesso file, nessun salvataggio: valori memorizzati e formule.
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        valori = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    fuori = [m[1] for w in caught if (m := re.search(r"Cell (\S+) is marked as a date", str(w.message)))]
    if fuori:
        valori.close()
        raise ConversionError(f"date fuori dall'intervallo di Excel in {', '.join(fuori)}: openpyxl le "
                              f"sostituirebbe con un errore, la conversione fedele non è possibile")
    formule = openpyxl.load_workbook(io.BytesIO(data), keep_links=False)
    try:
        sheets, n_formule, senza_valore = [], 0, []
        for wv, wf in zip(valori.worksheets, formule.worksheets):
            s = _Sheet(wv.title, _STATO.get(wv.sheet_state, ""))
            # ponytail: `_cells` = solo le celle presenti nel file. iter_rows scandirebbe l'intero
            # rettangolo (miliardi di celle con un solo valore in XFD1048576); API privata di
            # openpyxl 3.1, coperta da test_intervallo_enorme_non_scandisce_il_rettangolo.
            for (r, c), cell in wv._cells.items():
                if cell.value is not None:
                    s.put(r, c, cell.value, cell.number_format, error=cell.data_type == "e")
                comment, link = getattr(cell, "comment", None), getattr(cell, "hyperlink", None)
                if comment:
                    s.commento(r, c, comment.author, comment.text)
                if link and (link.target or link.location):
                    s.collegamento(r, c, link.target or "#" + link.location)
            con_formula = [k for k, cell in wf._cells.items() if cell.data_type == "f"]
            n_formule += len(con_formula)
            if mancanti := [k for k in con_formula if k not in s.celle]:
                senza_valore.append(f"{wv.title}!{_ranges(mancanti)}")
            s.unioni = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in wv.merged_cells.ranges]
            s.tabelle = sorted((t.displayName, t.ref) for t in wv.tables.values())
            sheets.append(s)
        grafici = len(valori.chartsheets)
    finally:
        valori.close()
        formule.close()
    avvisi = []
    if n_formule:
        avvisi.append(f"{n_formule} celle contengono formule: il Markdown riporta il valore memorizzato "
                      f"nel file, non ricalcolato; se il file è stato scritto da un programma e mai "
                      f"ricalcolato, può essere obsoleto.")
    if senza_valore:
        avvisi.append(f"Formule senza valore memorizzato, vuote nel Markdown: {'; '.join(senza_valore)}. "
                      f"Nessuna formula è stata ricalcolata.")
    with ZipFile(io.BytesIO(data)) as zf:
        oggetti = any(n.startswith(("xl/drawings/drawing", "xl/charts/", "xl/media/", "xl/embeddings/"))
                      for n in zf.namelist())
    if oggetti or grafici:
        avvisi.append("Il workbook contiene disegni, immagini, grafici o oggetti incorporati"
                      + (f" ({grafici} fogli grafico)" if grafici else "")
                      + ": non sono nel Markdown, restano nell'originale conservato in staging.")
    return sheets, avvisi


def _biff_counts(data: bytes) -> tuple[int, int, int]:
    """Formule, oggetti di disegno (esclusi i riquadri dei commenti) e grafici, contati sui
    record BIFF: xlrd legge i risultati delle formule ma non dice quali celle ne hanno."""
    import olefile  # dipendenza di markitdown[outlook], già usata per validare i binari

    with olefile.OleFileIO(io.BytesIO(data)) as ole:
        stream = ole.openstream("Workbook" if ole.exists("Workbook") else "Book").read()
    formule = oggetti = grafici = pos = 0
    while pos + 4 <= len(stream):
        rid, size = struct.unpack_from("<HH", stream, pos)
        if rid == 0x0006:                                            # FORMULA
            formule += 1
        elif rid == 0x005D and size >= 6:                            # OBJ: tipo in ftCmo
            oggetti += struct.unpack_from("<H", stream, pos + 8)[0] not in (0x05, 0x19)
        elif rid == 0x0809 and size >= 4:                            # BOF di un grafico
            grafici += struct.unpack_from("<H", stream, pos + 6)[0] == 0x0020
        pos += 4 + size
    return formule, oggetti, grafici


def _read_xls(data: bytes) -> tuple[list[_Sheet], list[str]]:
    import xlrd

    try:
        bk = xlrd.open_workbook(file_contents=data, formatting_info=True, ragged_rows=True,
                                logfile=io.StringIO())
    except Exception as exc:  # xlrd solleva tipi diversi a seconda del danno
        if "encrypted" in str(exc).lower():
            raise ConversionError("file protetto da password (cifrato): rimuovi la protezione "
                                  "e ricaricalo") from exc
        raise ConversionError(f"file danneggiato o non è un XLS leggibile ({type(exc).__name__}: "
                              f"{exc})") from exc
    try:
        epoch = MAC_EPOCH if bk.datemode else WINDOWS_EPOCH
        sheets = []
        for sh in bk.sheets():
            s = _Sheet(sh.name, _STATO.get(sh.visibility, ""))
            for r in range(sh.nrows):
                for c, cell in enumerate(sh.row(r)):
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    fmt = bk.format_map[bk.xf_list[cell.xf_index].format_key].format_str
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        s.put(r + 1, c + 1, xlrd.error_text_from_code[value], error=True)
                        continue
                    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype != xlrd.XL_CELL_TEXT and is_date_format(fmt):
                        # Stessa conversione di openpyxl per gli XLSX; fuori intervallo resta
                        # il seriale memorizzato, col formato elencato.
                        try:
                            value = from_excel(value, epoch, timedelta=is_timedelta_format(fmt))
                        except (OverflowError, ValueError):
                            pass
                    s.put(r + 1, c + 1, value, fmt)
            s.unioni = [(r0 + 1, c0 + 1, r1, c1) for r0, r1, c0, c1 in sh.merged_cells]
            for (r, c), note in sh.cell_note_map.items():
                s.commento(r + 1, c + 1, note.author, note.text)
            for (r, c), link in sh.hyperlink_map.items():
                if dest := link.url_or_path or (link.textmark and "#" + link.textmark):
                    s.collegamento(r + 1, c + 1, dest)
            sheets.append(s)
    finally:
        bk.release_resources()
    formule, oggetti, grafici = _biff_counts(data)
    avvisi = []
    if formule:
        avvisi.append(f"{formule} celle contengono formule: il Markdown riporta il valore memorizzato "
                      f"nel file, non ricalcolato. Dagli XLS xlrd non legge il testo delle formule.")
    if oggetti or grafici:
        avvisi.append(f"Il workbook contiene {oggetti} oggetti di disegno o immagini e {grafici} "
                      f"grafici: non sono nel Markdown, restano nell'originale conservato in staging.")
    return sheets, avvisi


def convert_excel(data: bytes, fmt: str) -> tuple[str, list[str]]:
    """Markdown e avvisi per la review. ConversionError per ogni lettura non fedele: nessun
    ripiego su convertitori che alterano le celle (pandas, normalizzazioni)."""
    motore, leggi = {"xlsx": ("openpyxl", _read_xlsx), "xls": ("xlrd", _read_xls)}[fmt]
    try:
        sheets, avvisi = leggi(data)
        if not any(s.celle for s in sheets):
            raise ConversionError("nessun valore nelle celle del workbook")
        return _render(sheets), avvisi
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"conversione Excel con {motore} fallita: "
                              f"{type(exc).__name__}: {exc}") from exc
