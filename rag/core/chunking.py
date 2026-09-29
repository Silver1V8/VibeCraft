"""Caricamento e chunking dei documenti: parsing del frontmatter YAML, loader multi-formato
(md/txt/docx/xlsx/pdf), estrazione delle entità, split per sezione (heading di livello 2,
le sezioni oltre MAX_CHUNK_CHARS in più parti) e archi tipizzati dal frontmatter. Nessuna dipendenza da Qdrant o dal modello di embedding."""
import re
import textwrap
from pathlib import Path

import networkx as nx
import yaml

from core.graph_ops import edge_trust_attrs


def split_frontmatter(content: str) -> tuple[str | None, str]:
    """Separa un documento in (testo del frontmatter, corpo). Il frontmatter esiste SOLO
    come blocco BEN FORMATO in cima al documento: una riga di sola apertura `---` e una
    riga di sola chiusura `---` (o `...`). Se manca la chiusura, non c'è frontmatter:
    `(None, contenuto intero)`.

    Il confronto è per RIGA e non per sottostringa. Un `split("---")` sull'intero testo
    (implementazione precedente) chiudeva il blocco alla prima occorrenza di `---` ovunque
    nel documento: se il frontmatter non era chiuso — tipico output di un LLM — il corpo
    Markdown finiva dentro il YAML, e la prima riga di elenco `- **Testo**: ...` faceva
    esplodere yaml.safe_load con un ScannerError su `*` (inizio di alias YAML)."""
    stripped = content.lstrip("﻿")  # tollera il BOM
    lines = stripped.splitlines(keepends=True)
    if not lines or lines[0].strip() != "---":
        return None, stripped
    for i in range(1, len(lines)):
        if lines[i].strip() in ("---", "..."):
            return "".join(lines[1:i]), "".join(lines[i + 1:])
    return None, stripped


def parse_markdown_file(filepath: str) -> tuple[dict, str]:
    """Estrae il Front Matter YAML e il corpo dal file Markdown. Un frontmatter che non
    sia un blocco delimitato in cima al file non viene interpretato (le righe orizzontali
    `---` a metà documento sono Markdown legittimo, non un delimitatore); un frontmatter
    che non produce un dict vale come assente, mai come metadata malformati."""
    with open(filepath, "r", encoding="utf-8") as f:
        content = f.read()
    fm_text, body = split_frontmatter(content)
    if fm_text is None:
        return {}, body
    metadata = yaml.safe_load(fm_text) or {}
    return (metadata if isinstance(metadata, dict) else {}), body


# I tre lettori qui sotto importano la propria libreria DENTRO la funzione (sono pesanti e
# servono solo per quel formato), ma NON la avvolgono piu' in un try/except che degrada a
# stringa vuota. Quel degrado sembrava tolleranza ed era perdita silenziosa di dati: con
# `pypdf` non dichiarata fra le dipendenze, ogni PDF in docs/ veniva indicizzato come
# documento VUOTO lasciando solo una riga di warning — indistinguibile, a valle, da un
# documento davvero senza testo (rag/AUDIT.md, A19). Ora le librerie sono tutte dichiarate
# in pyproject.toml: se manca, e' un ambiente rotto e deve fallire subito.
def _read_docx(filepath: str) -> str:
    import docx  # python-docx

    return "\n".join(p.text for p in docx.Document(filepath).paragraphs)


def _read_xlsx(filepath: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"# Foglio: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _read_pdf(filepath: str) -> str:
    import pypdf

    reader = pypdf.PdfReader(filepath)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def read_document(filepath_obj: Path) -> tuple[dict, str]:
    """Restituisce (frontmatter, corpo) per un documento. Il frontmatter YAML esiste
    solo nei .md; per gli altri formati è vuoto e si estrae solo il testo."""
    suffix = filepath_obj.suffix.lower()
    filepath = str(filepath_obj)
    if suffix == ".md":
        return parse_markdown_file(filepath)
    if suffix == ".txt":
        return {}, filepath_obj.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        return {}, _read_docx(filepath)
    if suffix == ".xlsx":
        return {}, _read_xlsx(filepath)
    if suffix == ".pdf":
        return {}, _read_pdf(filepath)
    return {}, ""

# Regioni da cui NON si estrae: dentro un blocco di codice le parentesi quadre sono
# sintassi (destrutturazioni `[k, v]`, literal `'[REDACTED]'`, classi di caratteri regex
# `[A-Z_]`), e nei link Markdown il testo fra quadre è un'etichetta di navigazione.
# Trattarli come riferimenti creava nodi-entità privi di semantica, che gonfiavano gli
# archi MENTIONS e sporcavano Semantic Map e query di impatto.
#
# NB: si scartano i soli blocchi delimitati da ```, NON il codice inline fra backtick
# singoli. La convenzione dei documenti è citare le regole proprio così — `[REGOLA-HITL-01]`
# nei titoli e nelle tabelle di riepilogo — quindi ignorare i backtick singoli perderebbe
# riferimenti veri (in vibe_coding_workflow le famiglie HITL e PROC compaiono SOLO in quella
# forma). Il rumore che si annidava nel codice inline lo scarta già il vaglio di forma sotto.
_CODE_FENCE_RE = re.compile(r"```.*?```", re.DOTALL)
_MD_LINK_RE = re.compile(r"\[[^\[\]]*\]\([^)\s]*\)")

# Un token candidato: contenuto di una coppia di quadre, non vuoto e non annidato.
# (Esclude di suo le checkbox Markdown `- [ ]` / `- [x]`, che non superano il vaglio sotto.)
_BRACKET_RE = re.compile(r"\[([^\[\]]+)\]")

# Identificatore di regola `[PREFISSO-CATEGORIA-NN]` conforme alla tassonomia della spec di
# ingestion ([REGOLA-VAL-07]). L'ULTIMO segmento deve essere numerico: è ciò che distingue
# una regola reale dai segnaposto dei template (`[REGOLA-XXX-NN]`, `[PREFISSO-CATEGORIA-XX]`,
# `[SAFEGUARD-YYY-NN]`), che sono esempi di formato e non devono entrare nel grafo.
_RULE_ID_RE = re.compile(r"^[A-Z][A-Z0-9]*(?:-[A-Z0-9]+)*-\d+$")

# Entità di dominio in forma `[UTENTE]` / `[AUDIT_LOG]`: sole maiuscole, cifre e underscore.
# Il trattino NON è ammesso qui, così le wildcard citate nella prosa (`[REGOLA-*]`,
# `[SAFEGUARD-*]`) e i frammenti tipo `A-Z_` restano fuori.
_ENTITY_ID_RE = re.compile(r"^[A-Z][A-Z0-9_]{2,}$")


def extract_entities(text: str) -> list[str]:
    """Entità e identificatori di regola citati nel corpo di un documento, nella forma
    `[UTENTE]` o `[REGOLA-VAL-07]`, in ordine di prima comparsa e senza duplicati.

    Prima dell'estrazione rimuove i blocchi di codice delimitati da ``` e i link Markdown:
    le parentesi quadre che compaiono lì sono sintassi o navigazione, mai riferimenti al
    dominio. Ogni candidato deve poi avere la forma di un identificatore — il che scarta le
    checkbox, i segnaposto dei template e le wildcard citate a scopo esplicativo."""
    for pattern in (_CODE_FENCE_RE, _MD_LINK_RE):
        text = pattern.sub(" ", text)
    found: list[str] = []
    seen: set[str] = set()
    for raw in _BRACKET_RE.findall(text):
        token = raw.strip()
        if token in seen:
            continue
        if _RULE_ID_RE.match(token) or _ENTITY_ID_RE.match(token):
            seen.add(token)
            found.append(token)
    return found


# Heading di livello 2 (## Sezione): unità naturale di chunking dei design contract.
_H2_RE = re.compile(r"(?m)^##\s+(.+?)\s*$")
_H3_RE = re.compile(r"^###\s")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]*-[\s:|-]*\|?\s*$")

# Tetto per chunk, sul testo completo (titolo + sezione + corpo). I consumatori troncano
# ogni chunk oltre il proprio limite — il grounding di doc_generation a 2500 caratteri,
# /ask a 3000 — perdendo in silenzio la coda della sezione: sull'indice v20260922T095110
# 24 chunk su 286 superavano 3000 caratteri, fino a 22.182 (gli scenari BDD di spec_auth,
# arrivati a /ask al 14%). Allineato al consumatore più stretto: una sezione più lunga
# diventa più parti, e nessuno taglia più. Cambiarlo richiede un full rebuild e la
# ricalibrazione congiunta di alpha, tetto del grafo e peso di sezione.
MAX_CHUNK_CHARS = 2500


def _slug(text: str, maxlen: int = 32) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return (s[:maxlen].strip("-") or "sez")


def _split_body(body: str, budget: int) -> list[str]:
    """Divide il corpo di una sezione troppo lunga in parti di al più `budget` caratteri.

    Taglia di preferenza prima di un `###` (se cade nella seconda metà della parte),
    altrimenti sull'ultima riga vuota fuori da un blocco di codice; solo se non ce n'è
    alcuna taglia sulla riga corrente. Una parte che prosegue un sottotitolo lo ripete con
    "(continua)": resta leggibile da sola, e retrieval.restates_only vede ancora il codice
    di regola DICHIARATO nel titolo, invece di scambiare la continuazione per una
    riformulazione da penalizzare. Un taglio forzato dentro un blocco ``` lo chiude e lo
    riapre; dentro una tabella ne ripete intestazione e separatore."""
    half = max(budget // 2, 1)
    lines = [piece for line in body.split("\n")
             for piece in (textwrap.wrap(line, half) if len(line) > half else [line])]
    parts: list[str] = []
    cur: list[str] = []
    size = head = 0        # caratteri di cur; righe di contesto ripetute in testa a cur
    cuts: list[tuple[int, int, str | None, bool]] = []  # (indice, caratteri prima, ### attivo, prima di un ###)
    h3 = fence = None
    table: list[str] = []  # intestazione e separatore della tabella in corso

    def cut(at: int, before: int, carry: list[str], close: str | None = None) -> None:
        nonlocal cur, size, head, cuts
        if any(x.strip() for x in cur[head:at]):
            parts.append("\n".join(cur[:at] + ([close] if close else [])).strip())
        shift = len(carry) - at
        extra = sum(len(x) + 1 for x in carry) - before
        cuts = [(i + shift, c + extra, h, b) for i, c, h, b in cuts if i > at]
        cur = carry + cur[at:]
        head = len(carry)
        size = sum(len(x) + 1 for x in cur)

    for line in lines:
        is_h3 = fence is None and bool(_H3_RE.match(line))
        is_row = fence is None and line.lstrip().startswith("|")
        if fence is None and (is_h3 or not line.strip()) and len(cur) > head:
            cuts.append((len(cur), size, h3, is_h3))
        while size + len(line) + 1 > budget and len(cur) > head:
            ok = [c for c in cuts if c[0] > head]
            if ok:
                at, before, h, before_h3 = max((c for c in ok if c[3] and c[1] >= size / 2),
                                               key=lambda c: c[0], default=ok[-1])
                cut(at, before, [f"{h} (continua)"] if h and not before_h3 else [])
            else:
                carry = [f"{h3} (continua)"] if h3 and not is_h3 else []
                if is_row and len(table) == 2 and _TABLE_SEP_RE.match(table[1]):
                    carry += table
                cut(len(cur), size, carry + ([fence] if fence else []),
                    close=fence[:3] if fence else None)
        cur.append(line)
        size += len(line) + 1
        if is_h3:
            h3 = line.strip()
        if _FENCE_RE.match(line):
            fence = None if fence else line.strip()
        table = (table + [line])[:2] if is_row else []
    if any(x.strip() for x in cur[head:]):
        parts.append("\n".join(cur).strip())
    return parts


def chunk_markdown(rel_path: str, metadata: dict, body: str) -> list[dict]:
    """Divide un documento Markdown in chunk per sezione (heading di livello 2),
    ancorando ogni chunk al titolo del documento (contesto del frontmatter).

    Restituisce una lista di dict {chunk_id, doc_id, section, text} con id stabili
    e leggibili nel formato `doc_id#NN-slug`. Se il documento non ha heading `##`,
    diventa un unico chunk. Una sezione oltre MAX_CHUNK_CHARS diventa più parti
    (vedi _split_body): la prima conserva l'id `doc_id#NN-slug`, le seguenti aggiungono
    `-pNN`; `section` resta il titolo della sezione, il testo dichiara "(parte i/n)"."""
    title = str(metadata.get("title") or rel_path)
    matches = list(_H2_RE.finditer(body))

    sections: list[tuple[str, str]] = []
    if not matches:
        sections.append(("(documento)", body))
    else:
        preamble = body[: matches[0].start()].strip()
        if preamble:
            sections.append(("(preambolo)", preamble))
        for i, m in enumerate(matches):
            start = m.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(body)
            sections.append((m.group(1).strip(), body[start:end].strip()))

    chunks: list[dict] = []
    for i, (heading, sec_body) in enumerate(sections):
        if not sec_body.strip():
            continue
        # Testo per l'embedding: ancorato al documento (titolo) e alla sezione,
        # così il chunk resta interpretabile fuori dal suo documento.
        header = f"# {title}\n## {heading}\n"
        bodies = [sec_body]
        if len(header) + len(sec_body) > MAX_CHUNK_CHARS:
            bodies = _split_body(sec_body, max(MAX_CHUNK_CHARS - len(header) - len(" (parte 99/99)"), 500))
        base = f"{rel_path}#{i:02d}-{_slug(heading)}"
        for j, part in enumerate(bodies, 1):
            label = heading if len(bodies) == 1 else f"{heading} (parte {j}/{len(bodies)})"
            chunks.append({
                "chunk_id": base if j == 1 else f"{base}-p{j:02d}",
                "doc_id": rel_path,
                "section": heading,
                "text": f"# {title}\n## {label}\n{part}",
            })

    if not chunks:  # fallback estremo: intero documento come singolo chunk
        chunks.append({"chunk_id": f"{rel_path}#00-doc", "doc_id": rel_path, "section": "(documento)", "text": body})
    return chunks

def add_frontmatter_edges(G: nx.DiGraph, doc_node_id: str, metadata: dict) -> int:
    """
    Indicizza le relazioni tipizzate dichiarate nel frontmatter YAML
    (graph_rag.edges: DERIVA_DA, VINCOLA, REFERENZIA, ...): archi semantici
    curati dall'autore del documento, molto più ricchi del semplice MENTIONS.
    """
    graph_rag = metadata.get("graph_rag") or {}
    edges = graph_rag.get("edges") or []
    added = 0
    for edge in edges:
        if not isinstance(edge, dict):
            continue
        raw_target = edge.get("target")
        if not raw_target:
            continue
        target = str(raw_target)
        relation = str(edge.get("type", "RELATES_TO"))
        if target not in G:
            G.add_node(target, type="entity")
        # curated: dichiarato a mano dall'autore nel frontmatter, non dedotto da un algoritmo
        # (una relazione non elencata in graph_trust.relation_nature ricade comunque qui).
        G.add_edge(doc_node_id, target, relation=relation, **edge_trust_attrs(relation))
        added += 1
    return added
