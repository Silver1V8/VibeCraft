"""Excel nell'intake: lettura diretta, XLSX con openpyxl e XLS con xlrd (core/intake_excel.py).

Le attese sono i valori SCRITTI nella fixture (archive/eval-results/excel_diretto_20260911/
crea_fixture.py, poi salvata da Excel 16 in XLS e XLSX) o costruiti nel test, mai l'uscita di un
altro convertitore. Il Markdown si legge con un parser CommonMark/GFM reale (markdown-it), come
lo leggerebbe chi lo usa: si verificano celle, coordinate e relazioni, non la presenza di parole.
"""
import hashlib
import io
from html.parser import HTMLParser
from pathlib import Path

import olefile
import openpyxl
import pytest
from markdown_it import MarkdownIt
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.table import Table

from core import intake_convert, intake_excel
from core.intake_convert import ConversionError, convert_to_markdown

from test_intake_staging import _client

FIXTURES = Path(__file__).parent / "fixtures" / "intake"


class _Html(HTMLParser):
    def __init__(self):
        super().__init__()
        self.sheet, self.tag, self.buf, self.tables, self.row = None, None, "", [], []

    def handle_starttag(self, tag, attrs):
        if tag in ("h2", "td", "th"):
            self.tag, self.buf = tag, ""
        elif tag == "br" and self.tag:
            self.buf += "\n"
        elif tag == "table":
            self.tables.append((self.sheet, []))
        elif tag == "tr":
            self.row = []

    def handle_data(self, data):
        if self.tag:
            self.buf += data

    def handle_endtag(self, tag):
        if tag == "h2":
            self.sheet, self.tag = self.buf, None
        elif tag in ("td", "th"):
            self.row.append(self.buf)
            self.tag = None
        elif tag == "tr":
            self.tables[-1][1].append(self.row)


def leggi(md: str):
    """{(foglio, riga, colonna): testo} delle celle non vuote e tabelle [(foglio, righe)]."""
    p = _Html()
    p.feed(MarkdownIt("commonmark").enable("table").render(md))
    celle = {}
    for sheet, (header, *body) in p.tables:
        assert header[0] == "Riga"
        cols = [column_index_from_string(h) for h in header[1:]]
        for row in body:
            assert len(row) == len(cols) + 1, row
            for col, text in zip(cols, row[1:]):
                key = (sheet, int(row[0]), col)
                assert key not in celle, f"cella duplicata {key}"
                if text:
                    celle[key] = text
    return celle, p.tables


def tabella_con(tables, sheet, riga):
    return next(rows for s, rows in tables if s == sheet and any(r[0] == str(riga) for r in rows[1:]))


def xlsx(build) -> bytes:
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Valori scritti da crea_fixture.py; Excel ha memorizzato i risultati delle formule della riga 7
# e degli errori della riga 9. Ogni altra cella deve essere vuota.
CASI = {
    ("Registro", 1, 1): "Registro casi",
    **{("Registro", 2, c): v for c, v in enumerate(["Codice", "Stato", "Quantità", "Data", "Importo", "Quota"], 1)},
    **{("Registro", 3, c): v for c, v in enumerate(["00123", "NA", "3", "2026-09-10", "1200.5", "0.12"], 1)},
    ("Registro", 4, 1): "00124", ("Registro", 4, 2): "N/A", ("Registro", 4, 4): "2026-09-11 00:00:00",
    ("Registro", 4, 5): "0", ("Registro", 4, 6): "FALSE",
    **{("Registro", 5, c): v for c, v in enumerate(["00125", "null", "7", "00:00:00", "False", "TRUE"], 1)},
    ("Registro", 7, 1): "Totale", ("Registro", 7, 3): "10", ("Registro", 7, 5): "1200.5", ("Registro", 7, 6): "2401",
    **{("Registro", 9, c): v for c, v in enumerate(["Errori", "#DIV/0!", "#N/A", "#N/A", "TRUE", "0"], 1)},
    **{("Registro", 10, c): v for c, v in enumerate(
        ["Durata", "36:00:00", "2026-09-10 13:45:30.250", "a|b", "riga1\nriga2",
         "*grassetto* _x_ codice_articolo <b>&amp;"], 1)},
    **{("Registro", 11, c): v for c, v in enumerate(
        ["  spazi  ", "\\percorso\\", "~~barrato~~", "`code`", "[link](http://x)", "=non formula"], 1)},
    **{("Affiancate", r, c): v for r, row in {1: ["Voce", "Euro", None, "Persona", "Anni"],
                                              2: ["A", "12", None, "Ada", "30"],
                                              3: ["B", "7", None, "Bruno", "41"],
                                              5: ["Nota", "Valore"], 6: ["x", "1"]}.items()
       for c, v in enumerate(row, 1) if v},
    ("Nascosto", 1, 1): "valore nascosto",
    ("Unite", 2, 2): "blocco", ("Unite", 2, 4): "accanto", ("Unite", 5, 1): "verticale",
}


@pytest.mark.parametrize("nome", ["casi_excel.xls", "casi_excel.xlsx"])
def test_fixture_excel_ogni_cella_alla_sua_coordinata_niente_di_piu(nome):
    r = convert_to_markdown(FIXTURES / nome, nome.rsplit(".", 1)[1])
    celle, tables = leggi(r.markdown)
    assert celle == CASI  # nessuna omissione, duplicazione o cella inventata
    assert r.metadata["motore"] == {"xls": "xlrd", "xlsx": "openpyxl"}[nome.rsplit(".", 1)[1]]
    assert r.metadata["fallback"] is None and "normalizzazione" not in r.metadata
    md = r.markdown
    assert [ln for ln in md.splitlines() if ln.startswith("## ")] == [
        "## Registro", "## Affiancate", "## Vuoto", "## Nascosto", "## Unite"]
    # Tipo ricostruibile: C9 è l'errore #N/A, D9 il TESTO «#N/A»; E9/F9 testo, E5 «False» no.
    assert "- Testo, non numero, data, booleano o errore: A3:A5, D9:F9." in md
    assert "- Celle unite, valore solo nella prima cella: A1:F1." in md
    assert "- Celle unite, valore solo nella prima cella: B2:C3, A5:A7." in md  # nessun valore duplicato
    assert "«00000» C5" in md and "«0%» F3" in md       # valore grezzo 7 e 0.12, formato elencato
    assert "- Foglio nascosto." in md and "- Nessuna cella con valore." in md
    assert "- Commento A7 (Revisore): Nota sul totale" in md
    assert "- Collegamento F2: https://example.org/quota" in md
    assert "codice_articolo" in md  # identificatore intero, cercabile senza escape
    avvisi = " ".join(r.metadata["avvisi"])
    assert "5 celle contengono formule" in avvisi and "non sono nel Markdown" in avvisi
    # La riga Totale resta una riga, allineata alle colonne delle intestazioni (A:F).
    assert tabella_con(tables, "Registro", 7) == [
        ["Riga", "A", "B", "C", "D", "E", "F"], ["7", "Totale", "", "10", "", "1200.5", "2401"]]
    assert tabella_con(tables, "Registro", 2)[0] == ["Riga", "A", "B", "C", "D", "E", "F"]


def test_stesso_workbook_xls_e_xlsx_stesso_markdown():
    assert (convert_to_markdown(FIXTURES / "casi_excel.xls", "xls").markdown
            == convert_to_markdown(FIXTURES / "casi_excel.xlsx", "xlsx").markdown)


@pytest.mark.parametrize("path,fmt,attesi", [
    (FIXTURES / "registro_excel.xlsx", "xlsx", ["Totale", "", "10", "", "1500.5", "2401"]),
    (FIXTURES / "registro.xls", "xls", ["Totale", "", "10", "", "1500.5", "2401"]),
    (FIXTURES / "legacy.xls", "xls", ["Totale", "", "10", "", "1500.5", "2401"]),
    # Scritto da openpyxl: C7 ha il risultato memorizzato OBSOLETO 999 (le celle sommate danno
    # 10), F7 nessun risultato. Nessun ricalcolo: si riporta ciò che il file contiene.
    (FIXTURES / "registro.xlsx", "xlsx", ["Totale", "", "999", "", "1500.5", ""]),
])
def test_riga_totale_con_celle_vuote_resta_associata(path, fmt, attesi):
    r = convert_to_markdown(path, fmt)
    _, tables = leggi(r.markdown)
    riga7 = tabella_con(tables, "Registro", 7)
    assert riga7 == [["Riga", "A", "B", "C", "D", "E", "F"], ["7", *attesi]]
    intestazioni = tabella_con(tables, "Registro", 2)
    assert intestazioni[0] == riga7[0]  # stesse colonne A:F: «10» sta sotto «Quantità»
    assert ["2", "Codice", "Stato", "Quantità", "Data", "Importo €", "Quota"] in intestazioni
    if attesi[-1] == "":
        assert any("Registro!F7" in a for a in r.metadata["avvisi"]), r.metadata["avvisi"]


def test_etichetta_e_valori_in_altra_posizione():
    def build(wb):
        ws = wb.active
        ws.title = "Quantità città"
        for r, values in [(4, ["ID", "Stato", "Pezzi", "Data", "Costo", "Quota"]),
                          (5, ["001", "NA", 3, None, 0, 0.2]), (8, [None, "Residuo Ω", None, 97, None, 1.5])]:
            for c, value in enumerate(values, 3):
                ws.cell(r, c, value)
    _, tables = leggi(convert_to_markdown(xlsx(build), "xlsx").markdown)
    assert tabella_con(tables, "Quantità città", 8) == [
        ["Riga", "C", "D", "E", "F", "G", "H"], ["8", "", "Residuo Ω", "", "97", "", "1.5"]]


@pytest.mark.parametrize("dichiarate", [False, True])
def test_tabelle_affiancate_e_sovrapposte_restano_distinte(dichiarate):
    def build(wb):
        ws = wb.active
        for row in (["Voce", "Euro", None, "Persona", "Anni"], ["A", 12, None, "Ada", 30], [],
                    ["Nota", "Valore"], ["x", 1]):
            ws.append(row)
        if dichiarate:
            ws.add_table(Table(displayName="Costi", ref="A1:B2"))
            ws.add_table(Table(displayName="Persone", ref="D1:E2"))
    md = convert_to_markdown(xlsx(build), "xlsx").markdown
    _, tables = leggi(md)
    assert [rows for _, rows in tables] == [
        [["Riga", "A", "B"], ["1", "Voce", "Euro"], ["2", "A", "12"]],
        [["Riga", "D", "E"], ["1", "Persona", "Anni"], ["2", "Ada", "30"]],
        [["Riga", "A", "B"], ["4", "Nota", "Valore"], ["5", "x", "1"]]]
    assert ("- Tabelle dichiarate nel file: Costi A1:B2; Persone D1:E2." in md) == dichiarate


def test_escape_markdown_e_testo_letterale():
    valori = ["a|b", "a\nb", "*x*", "_y_", "codice_articolo", "<script>alert(1)</script>", "&amp;",
              "c:\\dati\\", "`c`", "~~s~~", "[l](u)", "  NA  ", "\tt", "x" + chr(0x2028) + "y",
              "null", "N/A", "0012", "-3.5e2"]
    md = convert_to_markdown(xlsx(lambda wb: (setattr(wb.active, "title", "Budget # _x_"),
                                              wb.active.append(valori))), "xlsx").markdown
    celle, _ = leggi(md)
    assert [celle.get(("Budget # _x_", 1, c), "") for c in range(1, len(valori) + 1)] == valori
    assert "<script>" not in MarkdownIt("commonmark").enable("table").render(md)
    assert "- Testo, non numero, data, booleano o errore: Q1:R1." in md


def test_escape_di_cr_e_controlli():
    assert intake_excel._escape("a\r\nb") == "a&#13;<br>b"  # CRLF distinguibile da LF
    assert intake_excel._escape(" x" + chr(11)) == "&#32;x&#11;"


def test_formule_xlsx_senza_risultato_memorizzato_segnalate_non_ricalcolate():
    r = convert_to_markdown(xlsx(lambda wb: [wb.active.append(x) for x in
                                             (["Voce", "Importo"], ["A", 2], ["Totale", "=B2*3"])]), "xlsx")
    celle, _ = leggi(r.markdown)
    assert ("Sheet", 3, 2) not in celle  # nessun 6 inventato, nessuna formula spacciata per valore
    assert any(a.startswith("Formule senza valore memorizzato") and "Sheet!B3" in a
               for a in r.metadata["avvisi"]), r.metadata["avvisi"]


def test_intervallo_enorme_non_scandisce_il_rettangolo():
    def build(wb):
        ws = wb.active
        ws["A1"] = "inizio"
        ws["XFD1048576"] = "fine"
        ws["Z500"].font = Font(bold=True)  # formattazione residua senza valore
    md = convert_to_markdown(xlsx(build), "xlsx").markdown
    assert "### A1\n" in md and "### XFD1048576\n" in md and "Z500" not in md


def test_oltre_il_limite_errore_esplicito_mai_troncamento(monkeypatch):
    monkeypatch.setattr(intake_excel, "MAX_CELLS", 5)
    with pytest.raises(ConversionError, match="limite 5"):
        convert_to_markdown(xlsx(lambda wb: [wb.active.append([1, 2, 3]) for _ in range(2)]), "xlsx")


def test_troppe_celle_nel_file_rifiutate_prima_di_caricare(monkeypatch):
    """Anche le celle solo formattate contano: il limite protegge la memoria prima delle
    letture complete di openpyxl, con un errore esplicito."""
    monkeypatch.setattr(intake_excel, "MAX_XML_CELLS", 3)

    def build(wb):
        wb.active["A1"] = "unico valore"
        for ref in ("B2", "C3", "D4"):
            wb.active[ref].font = Font(bold=True)
    with pytest.raises(ConversionError, match="4 celle nei fogli"):
        convert_to_markdown(xlsx(build), "xlsx")


def test_workbook_senza_valori_errore():
    with pytest.raises(ConversionError, match="nessun valore"):
        convert_to_markdown(xlsx(lambda wb: wb.create_sheet("Altro")), "xlsx")


@pytest.mark.parametrize("refs", [("A1:XFD1048576",), ("A1:B3", "D1:E3")])
def test_unioni_troppo_estese_rifiutate_prima_di_openpyxl(monkeypatch, refs):
    """Il file piccolo non deve espandere milioni di MergedCell nel processo API."""
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx(lambda wb: wb.active.append(["x"])))) as src, \
            zipfile.ZipFile(buf, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                merges = "<mergeCells>" + "".join(f'<mergeCell ref="{ref}"/>' for ref in refs) + "</mergeCells>"
                data = data.replace(b"</worksheet>", merges.encode() + b"</worksheet>")
            dst.writestr(item, data)
    monkeypatch.setattr(intake_excel, "MAX_CELLS", 10)

    def vietato(*args, **kwargs):
        pytest.fail("openpyxl raggiunto prima del controllo sulle celle unite")

    monkeypatch.setattr(openpyxl, "load_workbook", vietato)
    with pytest.raises(ConversionError, match="celle unite.*limite"):
        convert_to_markdown(buf.getvalue(), "xlsx")


def test_xlsx_decompresso_troppo_grande_rifiutato_prima_di_leggere_xml(monkeypatch):
    import zipfile
    data = xlsx(lambda wb: wb.active.append(["testo"]))
    # Soglia piccola: include sharedStrings, immagini e parti senza elementi cella.
    monkeypatch.setattr(intake_excel, "MAX_XLSX_UNCOMPRESSED_BYTES", 1, raising=False)

    def vietato(*args, **kwargs):
        pytest.fail("contenuto ZIP letto prima del limite decompresso")

    monkeypatch.setattr(zipfile.ZipFile, "open", vietato)
    with pytest.raises(ConversionError, match="decompresso.*limite"):
        convert_to_markdown(data, "xlsx")


@pytest.mark.parametrize("nome", ["protetto.xls", "protetto.xlsx"])
def test_protetto_da_password_errore_chiaro(nome):
    with pytest.raises(ConversionError, match="protetto da password") as exc:
        convert_to_markdown(FIXTURES / nome, nome.rsplit(".", 1)[1])
    assert "danneggiato" not in str(exc.value)


def test_xls_con_stream_corrotto_errore_non_documento(tmp_path):
    path = tmp_path / "rotto.xls"
    path.write_bytes((FIXTURES / "legacy.xls").read_bytes())
    with olefile.OleFileIO(str(path), write_mode=True) as ole:
        ole.write_stream("Workbook", bytes(ole.get_size("Workbook")))
    with pytest.raises(ConversionError, match="danneggiato"):
        convert_to_markdown(path, "xls")


def test_xlsx_con_xml_corrotto_errore_senza_ripiego():
    buf = io.BytesIO()
    import zipfile
    with zipfile.ZipFile(io.BytesIO(xlsx(lambda wb: wb.active.append(["a"])))) as src, \
            zipfile.ZipFile(buf, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            dst.writestr(item, data[:40] if item.filename == "xl/worksheets/sheet1.xml" else data)
    with pytest.raises(ConversionError, match="openpyxl"):
        convert_to_markdown(buf.getvalue(), "xlsx")


def test_xls_non_richiede_libreoffice(monkeypatch):
    def vietato():
        raise AssertionError("LibreOffice non deve essere invocato per un XLS")
    monkeypatch.setattr(intake_convert, "_libreoffice_cmd", vietato)
    m = convert_to_markdown(FIXTURES / "legacy.xls", "xls").metadata
    assert (m["motore"], m["fallback"], m["versione_motore"]) == ("xlrd", None, "2.0.2")
    assert "normalizzazione" not in m


def test_conversione_deterministica():
    data = (FIXTURES / "casi_excel.xlsx").read_bytes()
    assert convert_to_markdown(data, "xlsx").markdown == convert_to_markdown(data, "xlsx").markdown


@pytest.mark.parametrize("nome", ["casi_excel.xlsx", "casi_excel.xls"])
def test_excel_api_preview_approve_solo_markdown_in_sandbox(nome, tmp_path, monkeypatch):
    import core.writer_lock as wl
    from core.chunking import parse_markdown_file
    monkeypatch.setattr(wl, "DATABASE_DIR", tmp_path)
    client = _client(tmp_path, monkeypatch)
    original = (FIXTURES / nome).read_bytes()
    upload = client.post("/api/intake/docs", files={"file": (nome, original)})
    assert upload.status_code == 200, upload.text
    doc_id = upload.json()["id"]
    preview = client.get(f"/api/intake/staged/{doc_id}/preview").json()
    assert leggi(preview["markdown"])[0] == CASI
    assert preview["metadata"]["motore"] == {"xls": "xlrd", "xlsx": "openpyxl"}[nome.rsplit(".", 1)[1]]
    assert preview["metadata"]["fallback"] is None
    assert preview["metadata"]["sha256"] == hashlib.sha256(original).hexdigest()
    assert not list((tmp_path / "docs").rglob("*"))
    approval = client.post(f"/api/intake/staged/{doc_id}/approve", json={
        "categoria": "04_knowledge_informative", "nome_file": nome})
    assert approval.status_code == 200, approval.text
    published = [p for p in (tmp_path / "docs").rglob("*") if p.is_file()]
    assert len(published) == 1 and published[0].name == nome + ".md"
    fm, body = parse_markdown_file(str(published[0]))
    assert fm["intake"]["motore"] == preview["metadata"]["motore"]
    assert fm["intake"]["sha256"] == hashlib.sha256(original).hexdigest()
    assert "normalizzazione" not in fm["intake"]
    assert leggi(body)[0] == CASI
    assert (tmp_path / "staging/docs" / preview["files"]["original"]).read_bytes() == original
