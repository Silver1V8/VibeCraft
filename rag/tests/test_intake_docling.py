"""Docling nell'intake (core/intake_convert.py): matrice dei motori, PDF digitale con titoli e
tabella, scansione con OCR locale, PDF danneggiato, fallback DICHIARATO quando Docling manca
o fallisce, esiti parziali, immagini mai pubblicate e flusso staging → docs/ con solo il .md.

I test con Docling reale richiedono i pesi in intake.docling.models_dir
(scripts/docling_models.py): senza, si saltano col motivo — oppure falliscono se
VIBECRAFT_REQUIRE_DOCLING=1, stesso patto di VIBECRAFT_REQUIRE_QDRANT. Gli altri sostituiscono
il solo convertitore Docling e girano ovunque."""
import hashlib
import importlib.metadata
import io
import os
import types
from pathlib import Path

import pytest

from core import intake, intake_convert
from core.chunking import chunk_markdown, parse_markdown_file
from core.config import INTAKE_ALLOWED_FORMATS
from core.intake_clean import clean_markdown, normalize_markdown
from core.intake_convert import ConversionError, convert_to_markdown

from test_intake_staging import _client, _journal_lines, _lock, _noop_reindex

FIXTURES = Path(__file__).parent / "fixtures" / "intake"
TABELLE = FIXTURES / "tabelle_titoli.pdf"  # titolo 24pt, sezioni 15/12pt, tabella 4x4, accenti
SAMPLE_PDF = FIXTURES / "sample.pdf"


@pytest.fixture
def staging_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(intake, "INTAKE_STAGING_DIR", tmp_path / "staging")
    return tmp_path / "staging"


@pytest.fixture
def docs_dir(tmp_path, monkeypatch):
    (tmp_path / "docs").mkdir()
    monkeypatch.setattr(intake, "DOCS_DIR", tmp_path / "docs")
    return tmp_path / "docs"


@pytest.fixture
def docling_reale(monkeypatch):
    mancanti = [rel for rel in {**intake_convert.DOCLING_WEIGHTS, **intake_convert.OCR_WEIGHTS}
                if not (intake_convert.INTAKE_DOCLING_MODELS_DIR / rel).is_file()]
    if mancanti:
        motivo = f"pesi Docling assenti in {intake_convert.INTAKE_DOCLING_MODELS_DIR}: {mancanti[0]}"
        if os.environ.get("VIBECRAFT_REQUIRE_DOCLING") == "1":
            pytest.fail(motivo)
        pytest.skip(motivo)
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)


@pytest.fixture
def senza_pesi(tmp_path, monkeypatch):
    """Docling attivo ma senza pesi: cartella vuota e nessun convertitore già in cache."""
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_MODELS_DIR", tmp_path / "modelli")
    monkeypatch.setattr(intake_convert, "_docling_converters", {})
    monkeypatch.setattr(intake_convert, "_weights_ok", set())
    return tmp_path / "modelli"


def _scansione(pdf: Path, dpi: int = 150) -> bytes:
    """Scansione sintetica: pagine renderizzate come immagini, nessun livello di testo."""
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(pdf.read_bytes())
    pagine = [doc[i].render(scale=dpi / 72).to_pil().convert("L") for i in range(len(doc))]
    buf = io.BytesIO()
    pagine[0].save(buf, format="PDF", save_all=True, append_images=pagine[1:], resolution=dpi)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Matrice dei motori
# ---------------------------------------------------------------------------
def test_matrice_docling_pdf_fra_i_formati_ammessi():
    """Esito del confronto per formato (NOTA_IMPLEMENTAZIONE_MIGLIORIE.md): Docling solo dove
    ha superato MarkItDown senza regressioni rilevate."""
    assert intake_convert.DOCLING_FORMATS == {"pdf"}
    assert intake_convert.DOCLING_FORMATS <= set(INTAKE_ALLOWED_FORMATS)


@pytest.mark.parametrize("formato,nome", [("txt", "sample.txt"), ("html", "sample.html"),
                                          ("docx", "sample.docx")])
def test_formati_non_coperti_restano_su_markitdown_senza_fallback(formato, nome, monkeypatch):
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    meta = convert_to_markdown(FIXTURES / nome, formato).metadata
    # Routing normale, non un ripiego: nessun fallback, nessun avviso.
    assert (meta["motore"], meta["fallback"], meta["avvisi"]) == ("markitdown", None, [])
    assert meta["versione_motore"] == importlib.metadata.version("markitdown")
    assert "pdf" not in meta


def test_docling_spento_i_pdf_tornano_a_markitdown(monkeypatch):
    """Rollback (intake.docling.enabled: false): il percorso storico, senza sonda né fallback."""
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", False)
    meta = convert_to_markdown(SAMPLE_PDF, "pdf").metadata
    assert (meta["motore"], meta["fallback"]) == ("markitdown", None)
    assert "pdf" not in meta


def test_opzioni_della_pipeline_sono_quelle_verificate(monkeypatch, tmp_path):
    from docling.datamodel.accelerator_options import AcceleratorDevice
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import TableFormerMode
    from docling.pipeline.standard_pdf_pipeline import StandardPdfPipeline
    monkeypatch.setattr(intake_convert, "_docling_converters", {})
    monkeypatch.setattr(intake_convert, "_verify_weights", lambda weights: None)

    for ocr in (False, True):
        fmt = intake_convert._docling_converter(ocr).format_to_options[InputFormat.PDF]
        o = fmt.pipeline_options
        assert fmt.pipeline_cls is StandardPdfPipeline
        assert o.accelerator_options.device == AcceleratorDevice.CPU
        assert o.enable_remote_services is False
        assert o.table_structure_options.mode == TableFormerMode.ACCURATE
        assert o.table_structure_options.do_cell_matching is True
        assert o.force_backend_text is True
        assert o.heading_hierarchy_options.enabled is True and o.generate_parsed_pages is True
        assert not (o.do_picture_description or o.do_picture_classification
                    or o.do_code_enrichment or o.do_formula_enrichment)
        assert o.artifacts_path == intake_convert.INTAKE_DOCLING_MODELS_DIR
        assert o.do_ocr is ocr and o.ocr_options.kind == "rapidocr" and o.ocr_options.lang == ["it"]


# ---------------------------------------------------------------------------
# Docling reale
# ---------------------------------------------------------------------------
def test_pdf_digitale_titoli_tabella_senza_ocr(docling_reale, tmp_path):
    r = convert_to_markdown(TABELLE, "pdf", filename="registro.pdf")
    meta, md = r.metadata, r.markdown

    assert meta["motore"] == "docling" and meta["versione_motore"] == importlib.metadata.version("docling")
    assert meta["fallback"] is None and meta["avvisi"] == []
    assert meta["pdf"] == {"pagine": 1, "pagine_senza_testo": 0, "ocr": False}
    headings = [ln for ln in md.splitlines() if ln.startswith("#")]
    assert headings == ["## Registro forniture", "### Regole di consegna", "### Tabella dei lotti",
                        "#### Note operative"]
    righe = [[c.strip() for c in ln.strip("|").split("|")] for ln in md.splitlines()
             if ln.startswith("|") and not ln.startswith("|--")]
    assert righe == [["Lotto", "Fornitore", "Quantità", "Stato"], ["L-01", "Acme Srl", "120", "consegnato"],
                     ["L-02", "Beta SpA", "45", "in attesa"], ["L-03", "Gamma Coop", "7", "annullato"]]
    # Identificatori e accenti come nella fonte: niente `\_`, niente segnaposto d'immagine.
    assert "identificativo_lotto" in md and "\\_" not in md and "<!--" not in md
    assert "la città di destinazione è indicata" in md
    assert chunk_markdown("registro.md", {"id": "registro"}, md)


def test_scansione_usa_ocr_locale_e_lo_dichiara(docling_reale):
    r = convert_to_markdown(_scansione(TABELLE), "pdf", filename="scansione.pdf")
    meta = r.metadata
    assert meta["motore"] == "docling" and meta["fallback"] is None
    assert meta["pdf"] == {"pagine": 1, "pagine_senza_testo": 1, "ocr": True}
    assert any("OCR locale" in a for a in meta["avvisi"])
    for parola in ("Registro forniture", "Beta SpA", "consegnato", "città"):
        assert parola in r.markdown


def test_pagina_bianca_errore_anche_con_ocr(docling_reale):
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument.new()
    doc.new_page(200, 200)
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(ConversionError, match="nessun testo estraibile"):
        convert_to_markdown(buf.getvalue(), "pdf", filename="bianca.pdf")


# ---------------------------------------------------------------------------
# Errori e fallback (senza pesi o con convertitore sostituito: girano ovunque)
# ---------------------------------------------------------------------------
def test_pdf_danneggiato_errore_chiaro(monkeypatch):
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    with pytest.raises(ConversionError, match="PDF illeggibile o danneggiato"):
        convert_to_markdown(b"%PDF-1.4\nnon un pdf vero", "pdf", filename="rotto.pdf")


def test_pdf_protetto_da_password_errore_esplicito(monkeypatch):
    from pypdf import PdfReader, PdfWriter
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    w = PdfWriter()
    w.add_page(PdfReader(TABELLE).pages[0])
    w.encrypt(user_password="segreta", owner_password="segreta")
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(ConversionError, match="protetto da password"):
        convert_to_markdown(buf.getvalue(), "pdf", filename="protetto.pdf")


def test_pesi_mancanti_fallback_dichiarato(senza_pesi):
    meta = convert_to_markdown(SAMPLE_PDF, "pdf").metadata
    assert meta["motore"] == "markitdown"
    assert meta["fallback"]["da"] == "docling" and "peso mancante" in meta["fallback"]["motivo"]
    assert meta["avvisi"][0].startswith("Docling non ha convertito questo file")


def test_peso_alterato_non_viene_usato(senza_pesi):
    rel = next(iter(intake_convert.DOCLING_WEIGHTS))
    (senza_pesi / rel).parent.mkdir(parents=True)
    (senza_pesi / rel).write_bytes(b"non il peso verificato")
    meta = convert_to_markdown(SAMPLE_PDF, "pdf").metadata
    assert meta["motore"] == "markitdown" and "SHA-256" in meta["fallback"]["motivo"]


def test_scansione_senza_docling_errore_esplicito(senza_pesi):
    """MarkItDown non fa OCR: senza Docling una scansione non diventa un documento vuoto."""
    with pytest.raises(ConversionError, match="nessun testo estraibile.*Docling non ha potuto"):
        convert_to_markdown(_scansione(TABELLE), "pdf", filename="scansione.pdf")


class _Doc:
    def __init__(self, md: str, pictures: int):
        self._md, self.pictures = md, [object()] * pictures

    def export_to_markdown(self, **_kw):
        return self._md


def _risultato(stato: str, md: str = "## Titolo\n\ntesto del documento", pictures: int = 0,
               timeout: bool = False):
    from docling.datamodel.base_models import ConversionStatus
    return types.SimpleNamespace(
        status=ConversionStatus(stato), document=_Doc(md, pictures),
        errors=[types.SimpleNamespace(error_message="pagina 3 non elaborata")],
        has_timeout_errors=lambda: timeout)


def _con_docling_finto(monkeypatch, risultato):
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    monkeypatch.setattr(intake_convert, "_docling_converter",
                        lambda ocr: types.SimpleNamespace(convert=lambda *a, **kw: risultato))


def test_timeout_risultato_parziale_dichiarato_incompleto(monkeypatch):
    _con_docling_finto(monkeypatch, _risultato("partial_success", timeout=True))
    meta = convert_to_markdown(SAMPLE_PDF, "pdf").metadata
    assert meta["motore"] == "docling"
    assert any("timeout" in a and "INCOMPLETO" in a for a in meta["avvisi"])


def test_timeout_prima_di_qualunque_testo_fallback_col_motivo(monkeypatch):
    """Misurato con Docling reale (timeout 0,05 s): nessun testo prodotto. Il ripiego su
    MarkItDown è dichiarato e il motivo nomina il timeout, non un generico «Markdown vuoto»."""
    _con_docling_finto(monkeypatch, _risultato("partial_success", md="", timeout=True))
    meta = convert_to_markdown(SAMPLE_PDF, "pdf").metadata
    assert meta["motore"] == "markitdown" and "timeout" in meta["fallback"]["motivo"]


def test_esito_parziale_senza_timeout_dichiarato(monkeypatch):
    _con_docling_finto(monkeypatch, _risultato("partial_success"))
    avvisi = convert_to_markdown(SAMPLE_PDF, "pdf").metadata["avvisi"]
    assert any("solo in parte" in a and "pagina 3 non elaborata" in a for a in avvisi)


def test_immagini_del_pdf_non_pubblicate_e_dichiarate(monkeypatch):
    _con_docling_finto(monkeypatch, _risultato("success", pictures=2))
    avvisi = convert_to_markdown(SAMPLE_PDF, "pdf").metadata["avvisi"]
    assert any(a.startswith("2 immagini o figure") for a in avvisi)


@pytest.mark.parametrize("risultato", [_risultato("failure"), _risultato("success", md="")])
def test_guasto_docling_fallback_markitdown_dichiarato(monkeypatch, risultato):
    """Esito `failure`, oppure Markdown vuoto da un PDF CON testo: guasto del motore, non del
    documento. Il PDF passa da MarkItDown e la preview lo dice."""
    _con_docling_finto(monkeypatch, risultato)
    r = convert_to_markdown(SAMPLE_PDF, "pdf")
    assert r.metadata["motore"] == "markitdown" and r.metadata["fallback"]["da"] == "docling"
    assert r.markdown.strip()


def test_scansione_vuota_anche_con_ocr_e_errore_non_fallback(monkeypatch):
    _con_docling_finto(monkeypatch, _risultato("success", md=""))
    with pytest.raises(ConversionError, match="l'OCR locale non ha riconosciuto testo"):
        convert_to_markdown(_scansione(SAMPLE_PDF), "pdf", filename="vuota.pdf")


def test_immagini_locali_rimosse_url_assoluti_conservati():
    html = (b"<html><body><h1>Titolo</h1><p>testo</p><img src=\"logo.png\" alt=\"logo\">"
            b"<img src=\"https://example.org/a.png\" alt=\"esterna\"></body></html>")
    r = convert_to_markdown(html, "html", filename="pagina.html")
    assert "logo.png" not in r.markdown and "https://example.org/a.png" in r.markdown
    assert any(a.startswith("1 riferimenti a immagini locali rimossi") for a in r.metadata["avvisi"])


def test_normalizzazione_docling_non_toglie_righe_di_contenuto():
    """Il cleaner delle impaginazioni toglierebbe un KPI isolato e una riga di tabella con i
    puntini: sul Markdown già strutturato di Docling si usa solo la normalizzazione."""
    md = "## Catalogo\n\n11\n\n| Voce | Pagina |\n|---|---|\n| Indice ..... | 4 |\n"
    assert "11" in normalize_markdown(md).splitlines() and "Indice ....." in normalize_markdown(md)
    assert "11" not in clean_markdown(md).splitlines() and "Indice ....." not in clean_markdown(md)


# ---------------------------------------------------------------------------
# Flusso applicativo: staging → preview → approve/reject, solo .md in docs/
# ---------------------------------------------------------------------------
def test_pdf_docling_da_staging_a_docs_solo_md(docling_reale, staging_dir, docs_dir, tmp_path):
    pdf = TABELLE.read_bytes()
    record = intake.stage_document("registro.pdf", pdf)
    assert list(docs_dir.rglob("*")) == []  # nulla in docs/ prima dell'approvazione

    preview = intake.get_preview(record["id"])
    assert preview["metadata"]["motore"] == "docling" and preview["metadata"]["avvisi"] == []
    assert "Beta SpA" in preview["markdown"]

    intake.approve(record["id"], "04_knowledge_informative", "registro-forniture",
                   reindex=_noop_reindex, lock_path=_lock(tmp_path))
    pubblicati = [p for p in docs_dir.rglob("*") if p.is_file()]
    assert [p.name for p in pubblicati] == ["registro-forniture.md"]
    fm, body = parse_markdown_file(str(pubblicati[0]))
    sha = hashlib.sha256(pdf).hexdigest()
    assert fm["intake"]["motore"] == "docling" and fm["intake"]["sha256"] == sha
    assert fm["intake"]["versione_motore"] == importlib.metadata.version("docling")
    assert "![" not in body  # nessun riferimento ad asset non pubblicati
    # L'originale resta recuperabile in staging, identico byte per byte.
    originale = staging_dir / "docs" / record["files"]["original"]
    assert hashlib.sha256(originale.read_bytes()).hexdigest() == sha


def test_scansione_da_staging_a_docs_con_avviso_ocr(docling_reale, staging_dir, docs_dir, tmp_path):
    record = intake.stage_document("scansione.pdf", _scansione(TABELLE))
    assert record["metadata"]["pdf"]["ocr"] is True
    intake.approve(record["id"], "04_knowledge_informative", "scansione",
                   reindex=_noop_reindex, lock_path=_lock(tmp_path))
    assert [p.suffix for p in docs_dir.rglob("*") if p.is_file()] == [".md"]


def test_rifiuto_non_pubblica_e_conserva_originale(staging_dir, docs_dir):
    record = intake.stage_document("nota.txt", b"testo da rifiutare")
    intake.reject(record["id"], "prova")
    assert list(docs_dir.rglob("*")) == []
    originale = staging_dir / "rejected" / record["files"]["original"]
    assert originale.read_bytes() == b"testo da rifiutare"


def test_errore_di_conversione_non_lascia_nulla(staging_dir, docs_dir, monkeypatch):
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    with pytest.raises(ConversionError):
        intake.stage_document("rotto.pdf", b"%PDF-1.4\nnon un pdf vero")
    assert list((staging_dir / "docs").iterdir()) == [] and _journal_lines(staging_dir) == []
    assert list(docs_dir.rglob("*")) == []


@pytest.mark.parametrize("nome,contenuto,status", [
    ("rotto.pdf", b"%PDF-1.4\nnon un pdf vero", 422),
    ("vuoto.pdf", b"", 422),
    ("script.exe", b"MZ", 415),
])
def test_route_upload_errori_non_pubblicano_nulla(tmp_path, monkeypatch, nome, contenuto, status, legacy_workflow_project):
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    client = _client(tmp_path, monkeypatch)
    resp = client.post("/api/intake/docs", files={"file": (nome, contenuto, "application/octet-stream")})
    assert resp.status_code == status, resp.text
    assert client.get("/api/intake/staged").json()["staged"] == []
    assert list((tmp_path / "docs").rglob("*")) == []


def test_route_preview_espone_motore_e_avvisi_del_fallback(tmp_path, monkeypatch, senza_pesi, legacy_workflow_project):
    client = _client(tmp_path, monkeypatch)
    doc_id = client.post("/api/intake/docs", files={
        "file": ("guida.pdf", SAMPLE_PDF.read_bytes(), "application/pdf")}).json()["id"]
    meta = client.get(f"/api/intake/staged/{doc_id}/preview").json()["metadata"]
    assert meta["motore"] == "markitdown" and meta["fallback"]["da"] == "docling"
    assert meta["avvisi"] and meta["pdf"]["pagine"] == 1


# ---------------------------------------------------------------------------
# XLSX: lettura diretta con openpyxl (i casi completi sono in test_intake_excel.py)
# ---------------------------------------------------------------------------
def _xlsx(fogli: dict[str, list[list]]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome, righe in fogli.items():
        ws = wb.create_sheet(nome)
        for riga in righe:
            ws.append(riga)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _righe(md: str) -> list[list[str]]:
    return [[c.strip() for c in ln.strip()[1:-1].split("|")][1:] for ln in md.splitlines()
            if ln.lstrip().startswith("|") and not set(ln.strip()) <= set("|-: ")]


def test_xlsx_nomi_dei_fogli_e_valori_letterali(monkeypatch):
    """MarkItDown (pandas) trasformava in vuote le celle «NA», «N/A», «null», scriveva 7.0
    per 7 e «NaN» per le vuote: dati alterati. Openpyxl conserva le celle con coordinate esplicite."""
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    data = _xlsx({"Registro": [["Codice", "Stato", "Quantità"], ["00123", "NA", 3],
                               ["00124", "N/A", None], ["00125", "null", 7]],
                  "Contatti città": [["Nome", "Città"], ["Anna", "Forlì"]]})
    r = convert_to_markdown(data, "xlsx", filename="registro.xlsx")
    assert (r.metadata["motore"], r.metadata["fallback"]) == ("openpyxl", None)
    titoli = [ln for ln in r.markdown.splitlines() if ln.startswith("## ")]
    assert titoli == ["## Registro", "## Contatti città"]
    righe = _righe(r.markdown)
    assert ["00123", "NA", "3"] in righe and ["00125", "null", "7"] in righe
    assert ["00124", "N/A", ""] in righe and ["Anna", "Forlì"] in righe
    assert "NaN" not in r.markdown


def test_xlsx_formula_mai_calcolata_resta_vuota(monkeypatch):
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    data = _xlsx({"Budget": [["Voce", "Importo"], ["A", 2], ["Totale", "=B2*3"]]})
    r = convert_to_markdown(data, "xlsx", filename="b.xlsx")
    assert ["Totale", ""] in _righe(r.markdown)  # nessun 6 inventato, nessuna formula spacciata per valore
    # ...ma la cella vuota non passa in silenzio: la review deve sapere che manca un risultato.
    assert any(a.startswith("Formule senza valore memorizzato") and "Budget!B3" in a
               for a in r.metadata["avvisi"]), r.metadata["avvisi"]


def test_xlsx_formula_con_valore_memorizzato_nessun_avviso(monkeypatch):
    import zipfile as zf
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", True)
    data = _xlsx({"Budget": [["Voce", "Importo"], ["A", 2], ["Totale", "=B2*3"]]})
    src, out = zf.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zf.ZipFile(out, "w") as dst:  # risultato memorizzato, come dopo un salvataggio di Excel
        for item in src.infolist():
            xml = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = xml.replace(b"<f>B2*3</f><v></v>", b"<f>B2*3</f><v>6</v>").replace(
                    b"<f>B2*3</f><v />", b"<f>B2*3</f><v>6</v>").replace(b"<f>B2*3</f></c>", b"<f>B2*3</f><v>6</v></c>")
                assert b"<v>6</v>" in xml
            dst.writestr(item, xml)
    r = convert_to_markdown(out.getvalue(), "xlsx", filename="b.xlsx")
    assert ["Totale", "6"] in _righe(r.markdown)
    assert not any("formula" in a for a in r.metadata["avvisi"])


def test_xlsx_guasto_openpyxl_errore_senza_fallback_lossy(monkeypatch):
    import openpyxl
    data = _xlsx({"F": [["a", "b"], [1, 2]]})

    def _guasto(*_a, **_kw):
        raise RuntimeError("guasto simulato")
    monkeypatch.setattr(openpyxl, "load_workbook", _guasto)
    with pytest.raises(ConversionError, match="openpyxl.*guasto simulato"):
        convert_to_markdown(data, "xlsx", filename="f.xlsx")


def test_xlsx_indipendente_dal_flag_pdf_docling(monkeypatch):
    monkeypatch.setattr(intake_convert, "INTAKE_DOCLING_ENABLED", False)
    meta = convert_to_markdown(_xlsx({"F": [["a", "b"], [1, 2]]}), "xlsx", filename="f.xlsx").metadata
    assert (meta["motore"], meta["fallback"]) == ("openpyxl", None)


# ---------------------------------------------------------------------------
# File problematici e asset
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("formato", ["xlsx", "docx", "pptx"])
def test_ooxml_protetto_da_password_errore_esplicito_non_danneggiato(formato):
    """Fixture reale: XLSX salvato da Excel 16 con password (contenitore OLE con
    EncryptedPackage). Prima: «file danneggiato o non è un XLSX valido»."""
    data = (FIXTURES / "protetto.xlsx").read_bytes()
    with pytest.raises(ConversionError, match="protetto da password") as exc:
        convert_to_markdown(data, formato, filename=f"riservato.{formato}")
    assert "danneggiato" not in str(exc.value)


@pytest.mark.parametrize("formato", ["docx", "pptx", "xlsx", "epub", "xls", "msg"])
def test_file_danneggiato_errore_non_documento(formato):
    """Prima un DOCX rotto diventava un «documento» fatto dei suoi byte grezzi."""
    with pytest.raises(ConversionError, match="danneggiato"):
        convert_to_markdown(b"PK\x03\x04 byte che non sono un contenitore", formato,
                            filename=f"rotto.{formato}")


def test_txt_sintassi_immagine_resta_testo_della_fonte():
    r = convert_to_markdown(b"Esempio: ![logo](logo.png) va scritto cosi'.", "txt", filename="n.txt")
    assert "![logo](logo.png)" in r.markdown and r.metadata["avvisi"] == []


def test_immagine_locale_in_blocco_di_codice_resta():
    html = b"<html><body><h1>T</h1><pre><code>![x](locale.png)</code></pre></body></html>"
    assert "![x](locale.png)" in convert_to_markdown(html, "html", filename="c.html").markdown


def test_immagini_incorporate_scartate_da_markitdown_sono_dichiarate():
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    d = Document()
    d.add_heading("Con immagine", 1)
    d.add_paragraph("testo")
    img = io.BytesIO()
    Image.new("RGB", (20, 10), "red").save(img, format="PNG")
    d.add_picture(io.BytesIO(img.getvalue()), width=Inches(1))
    buf = io.BytesIO()
    d.save(buf)
    r = convert_to_markdown(buf.getvalue(), "docx", filename="img.docx")
    assert "data:image" not in r.markdown
    assert any(a.startswith("1 immagini incorporate") for a in r.metadata["avvisi"])


def test_titolo_del_frontmatter_dal_titolo_piu_alto(staging_dir, docs_dir, tmp_path):
    record = intake.stage_document("nota.md", b"## Titolo vero\n\n### Dettaglio\n\ntesto\n")
    intake.approve(record["id"], "04_knowledge_informative", "nota-x",
                   reindex=_noop_reindex, lock_path=_lock(tmp_path))
    fm, _body = parse_markdown_file(str(docs_dir / "04_knowledge_informative" / "nota-x.md"))
    assert fm["title"] == "Titolo vero"
